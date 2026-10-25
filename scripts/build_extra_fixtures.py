"""PPTX + ZIP fixture builders for the representative spike."""
from __future__ import annotations

import io
import zipfile
from pathlib import Path


def build_pptx(path: str | Path = "data/fixtures/spike/slides.pptx") -> Path:
    from pptx import Presentation
    from pptx.util import Inches
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    slide.shapes.title.text = "Maths Warm-up"
    slide.shapes.add_textbox(Inches(0.5), Inches(1.5), Inches(9), Inches(3)).text_frame.text = (
        "1. What is 7 x 8?\nA) 54 B) 56 C) 58\n\n2. Simplify 12/18.\nA) 1/2 B) 2/3")
    prs.save(str(path))
    return path


def build_zip(path: str | Path = "data/fixtures/spike/bundle.zip") -> Path:
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as zf:
        zf.writestr("worksheet_extra.txt",
                    "1. What is 5 + 5?\nA) 9 B) 10\n2. What is 9 - 4?\nA) 5 B) 6\n")
        zf.writestr("answers_extra.txt", "Answer Key\n1. B\n2. A\n")
    path.write_bytes(buf.getvalue())
    return path


if __name__ == "__main__":
    print(f"wrote {build_pptx()}")
    print(f"wrote {build_zip()}")
