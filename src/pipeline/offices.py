"""Step: PPTX + TXT extraction. PPTX slides become SourcePages;
TXT files become single pages. Keeps the spike's 11-type coverage."""
from __future__ import annotations

from pathlib import Path

from sqlalchemy.orm import Session

from src.core.database import SourceDocument, SourcePage
from src.core.logging import get_logger
from src.pipeline.extract import _ensure_pages

log = get_logger("pipeline.offices")


def run_office_extract(session: Session, docs: list[SourceDocument],
                       source_root: str | Path) -> dict:
    root = Path(source_root)
    done = 0
    for doc in docs:
        target = root / doc.source_path
        if not target.exists():
            continue
        if doc.file_type == "pptx":
            try:
                from pptx import Presentation
                prs = Presentation(str(target))
                pages = _ensure_pages(session, doc, len(prs.slides))
                for page, slide in zip(pages, prs.slides):
                    texts = [s.text for s in slide.shapes if hasattr(s, "text")]
                    joined = "\n".join(texts)
                    if len(joined.strip()) > len(page.text_content.strip()):
                        page.text_content = joined
                        page.text_confidence = 0.85
                    done += 1
            except Exception as exc:
                log.warning("pptx extract failed for %s: %s", doc.source_path, exc)
        elif doc.file_type == "text":
            pages = _ensure_pages(session, doc, 1)
            try:
                text = target.read_text(encoding="utf-8", errors="replace")
                if len(text.strip()) > len(pages[0].text_content.strip()):
                    pages[0].text_content = text
                    pages[0].text_confidence = 1.0
                done += 1
            except Exception as exc:
                log.warning("text extract failed for %s: %s", doc.source_path, exc)
    session.commit()
    return {"pages": done}
